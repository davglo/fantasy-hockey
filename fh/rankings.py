"""Dave's rankings workbook: projections (FP) per player + scoring-settings check vs ESPN."""
from __future__ import annotations

import logging
import re
import unicodedata
from pathlib import Path

from fh import config

log = logging.getLogger(__name__)

# Settings-sheet label -> ESPN stat ID. Labels the league can't score (GAA, SV%, FO%, TOI...) are omitted.
LABEL_TO_STAT = {
    "goals": 13, "assists": 14, "points": 16, "shots": 29, "blocks": 32, "hits": 31,
    "powerplay goals": 18, "powerplay points": 38, "shorthanded goals": 20, "shorthanded points": 39,
    "plus-minus": 15, "penalty minutes": 17, "game-winning goals": 22, "faceoff wins": 23,
    "faceoff losses": 24, "wins": 1, "losses": 2, "overtime losses": 9, "shutouts": 7, "saves": 6,
    "goals against": 4,
}

NAME_ALIASES = {  # sheet spelling -> ESPN spelling (normalized); extend as mismatches show up
}


def norm_name(name: str) -> str:
    s = unicodedata.normalize("NFKD", name).encode("ascii", "ignore").decode().lower()
    s = re.sub(r"\b(jr|sr|ii|iii)\b", "", s)
    s = re.sub(r"[^a-z]", "", s)
    return NAME_ALIASES.get(s, s)


def initial_key(name: str) -> str:
    """First initial + last name, for 'Alex' vs 'Alexander' style mismatches."""
    parts = name.split()
    return norm_name(parts[0])[:1] + norm_name(" ".join(parts[1:])) if len(parts) > 1 else norm_name(name)


def norm_team(abbrev: str | None) -> str:
    t = (abbrev or "").replace(".", "").upper()
    return {"UTAH": "UTA", "VEG": "VGK", "LVK": "VGK", "MON": "MTL", "NAS": "NSH", "CLS": "CBJ",
            "WAS": "WSH", "CAL": "CGY", "NJD": "NJ", "SJS": "SJ", "TBL": "TB", "LAK": "LA"}.get(t, t)


def _open(path: Path):
    import openpyxl  # imported lazily so a missing package degrades to "no sheet"
    return openpyxl.load_workbook(path, read_only=True, data_only=True)


def load(path: Path = config.RANKINGS_XLSX) -> list:
    """Rows from 'The List': name, team, pos, fp, gp, sheet_rank, sheet_adp, boost."""
    if not path.exists():
        log.warning("rankings workbook not found at %s; using ESPN projections only", path)
        return []
    try:
        ws = _open(path)["The List"]
    except Exception as e:  # noqa: BLE001 - any workbook problem -> degrade, don't crash
        log.warning("could not read rankings workbook (%s); using ESPN projections only", e)
        return []
    rows = ws.iter_rows(values_only=True)
    header = next(rows)
    col = {}
    for i, h in enumerate(header):
        if h and h not in col:
            col[h] = i
    out = []
    for r in rows:
        name, fp = r[col["NAME"]], r[col["FP"]]
        if not name or not isinstance(fp, (int, float)):
            continue
        is_g = r[col["POS"]] == "G"
        gp = r[35] if is_g else r[col["GP"]]  # goalie GP lives in the second GP column
        out.append({
            "name": name, "key": norm_name(name), "team": norm_team(r[col["TEAM"]]), "pos": r[col["POS"]],
            "fp": float(fp), "gp": gp if isinstance(gp, (int, float)) else None,
            "sheet_rank": r[col["RK"]], "sheet_adp": r[col["ADP"]] if isinstance(r[col["ADP"]], (int, float)) else None,
            "boost": r[col["ADJ"]] if r[col["ADJ"]] not in (0, None) else "",
        })
    if not out:
        log.warning("rankings workbook has no cached FP values - open it in Excel and save once")
    return out


def check_scoring(league_scoring: dict, path: Path = config.RANKINGS_XLSX) -> list:
    """Compare the workbook's Points column with ESPN league scoring. Returns human-readable issues."""
    if not path.exists():
        return ["rankings workbook not found"]
    try:
        ws = _open(path)["Settings"]
    except Exception as e:  # noqa: BLE001
        return ["could not read Settings sheet: %s" % e]
    sheet = {}
    for label, pts, *_ in ws.iter_rows(min_col=1, max_col=2, values_only=True):
        if isinstance(label, str) and label.strip().lower() in LABEL_TO_STAT:
            sheet[LABEL_TO_STAT[label.strip().lower()]] = float(pts) if isinstance(pts, (int, float)) else 0.0
    issues = []
    for stat in sorted(set(sheet) | set(league_scoring)):
        want, have = float(league_scoring.get(stat, 0)), sheet.get(stat)
        name = config.STAT_NAMES.get(stat, str(stat))
        if have is None:
            if want:
                issues.append("%s: league scores %g, workbook can't model it" % (name, want))
        elif abs(want - have) > 1e-9:
            issues.append("%s: league %g vs workbook %g" % (name, want, have))
    return issues
